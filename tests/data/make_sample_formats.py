#!/usr/bin/env python3
"""Writes the books in tests/data/formats/ that the host tests convert.

One small book, "Formats Test" by Ona Rašytoja, in the formats Vellum turns
into HTML when it first sees them: two chapters, the second with a
sub-heading, bold and italic words, Lithuanian letters, a picture (top half
black) and a cover (bottom half black).

  sample.mobi   old Kindle format (MOBI 6, PalmDOC), made by Calibre
  sample.azw3   new Kindle format (KF8), made by Calibre
  sample.docx   Word, made with python-docx
  sample.rtf    rich text, written here (Windows-1257 and \\u escapes)
  sample.pdf    PDF with an outline, a page number and a running head on
                every page, made with ReportLab (ASCII85 + Flate streams)
  locked.pdf    sample.pdf with a password (pikepdf)
  scan.pdf      one page that is only a picture
  sample.cbz    a comic: two folders, pages named page1, page2, page10 (one
                PNG, stored and deflated entries) and a ComicInfo.xml

Needs Calibre's ebook-convert, Pillow, python-docx, reportlab and pikepdf.
"""
import io
import pathlib
import subprocess
import tempfile
import zipfile

import docx
import pikepdf
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A6
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

out = pathlib.Path(__file__).resolve().parent / "formats"
out.mkdir(exist_ok=True)


def half(w, h, top_black, fmt):
    img = Image.new("L", (w, h), 255)
    box = [0, 0, w, h // 2] if top_black else [0, h // 2, w, h]
    ImageDraw.Draw(img).rectangle(box, fill=0)
    buf = io.BytesIO()
    img.convert("RGB" if fmt == "JPEG" else "L").save(buf, fmt)
    return buf.getvalue()


PICTURE = half(200, 120, True, "PNG")
COVER = half(300, 450, False, "JPEG")

HTML = """<html lang="lt"><head><meta charset="utf-8"><title>Formats Test</title>
<meta name="author" content="Ona Rašytoja"></head><body>
<h1>Chapter One</h1>
<p>The first paragraph has <b>bold</b> and <i>italic</i> words.</p>
<p><img src="pic.png" alt="Half black"/></p>
<p>After the picture.</p>
<h1>Chapter Two</h1>
<h2>A Part</h2>
<p>Ąžuolas ėjo per šilą.</p>
</body></html>
"""

with tempfile.TemporaryDirectory() as tmp:
    t = pathlib.Path(tmp)
    (t / "pic.png").write_bytes(PICTURE)
    (t / "cover.jpg").write_bytes(COVER)
    (t / "book.html").write_text(HTML, encoding="utf-8")
    common = ["--cover", str(t / "cover.jpg"), "--language", "lt", "--authors", "Ona Rašytoja"]
    subprocess.run(["ebook-convert", str(t / "book.html"), str(out / "sample.mobi"), "--mobi-file-type", "old"] + common,
                   check=True, capture_output=True)
    subprocess.run(["ebook-convert", str(t / "book.html"), str(out / "sample.azw3")] + common,
                   check=True, capture_output=True)

# Word.
d = docx.Document()
d.core_properties.title = "Formats Test"
d.core_properties.author = "Ona Rašytoja"
d.core_properties.language = "lt-LT"
d.add_heading("Chapter One", level=1)
p = d.add_paragraph("The first paragraph has ")
p.add_run("bold").bold = True
p.add_run(" and ")
p.add_run("italic").italic = True
p.add_run(" words.")
d.add_picture(io.BytesIO(PICTURE))
d.add_paragraph("After the picture.")
d.add_heading("Chapter Two", level=1)
d.add_heading("A Part", level=2)
d.add_paragraph("Ąžuolas ėjo per šilą.")
d.save(out / "sample.docx")

# Rich text: the title in \info, headings by outline level and style name,
# Lithuanian letters in the code page and as \u escapes.
RTF = (r"{\rtf1\ansi\ansicpg1257\deff0{\fonttbl{\f0 Times;}}"
       r"{\stylesheet{\s0 Normal;}{\s1\b\fs32 heading 1;}{\s2\b\fs28 heading 2;}}"
       r"{\info{\title Formats Test}{\author Ona Ra\'f0ytoja}}"
       r"{\*\generator Test;}"
       "\n"
       r"{\pard\s1 Chapter One\par}"
       r"{\pard The first paragraph has {\b bold} and {\i italic} words.\par}"
       r"{\pard After the picture.\par}"
       r"{\pard\s1 Chapter Two\par}"
       r"{\pard\outlinelevel1 A Part\par}"
       r"{\pard \u260?\'feuolas \u279?jo per \'f0il\u261?.\par}"
       "}")
(out / "sample.rtf").write_bytes(RTF.encode("latin-1"))

# PDF with an outline, page numbers and a running head.
styles = getSampleStyleSheet()
body = styles["BodyText"]
body.fontName = "Times-Roman"


class Doc(SimpleDocTemplate):
    def afterFlowable(self, f):
        if isinstance(f, Paragraph) and f.style.name in ("Heading1", "Heading2"):
            key = "k%d" % id(f)
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(f.getPlainText(), key, level=0 if f.style.name == "Heading1" else 1)


def on_page(c, doc):
    c.setFont("Helvetica", 7)
    c.drawCentredString(A6[0] / 2, 14, str(doc.page))
    c.drawCentredString(A6[0] / 2, A6[1] - 18, "Formats Test")


words = ("It was a bright cold day in April, and the clocks were striking thirteen. "
         "Nobody in the house had heard them, and the cat went on sleeping. ")
story = [Paragraph("Chapter One", styles["Heading1"]),
         Paragraph("The first paragraph has <b>bold</b> and <i>italic</i> words. " + words * 3, body),
         Paragraph("The second paragraph. " + words * 4, body),
         PageBreak(),
         Paragraph("Chapter Two", styles["Heading1"]),
         Paragraph("Second chapter text. " + words * 2, body),
         Paragraph("A Part", styles["Heading2"]),
         Paragraph("Short.", body)]
Doc(str(out / "sample.pdf"), pagesize=A6, title="Formats Test", author="Ona Rasytoja").build(
    story, onFirstPage=on_page, onLaterPages=on_page)
with pikepdf.open(out / "sample.pdf") as pdf:
    pdf.save(out / "locked.pdf", encryption=pikepdf.Encryption(owner="o", user="u", R=4))
Image.open(io.BytesIO(PICTURE)).convert("RGB").save(out / "scan.pdf")

# A comic.
with zipfile.ZipFile(out / "sample.cbz", "w") as z:
    for i, name in enumerate(["Vol 1/page10.jpg", "Vol 1/page2.jpg", "Vol 1/page1.png", "Vol 2/p1.jpg"]):
        fmt = "PNG" if name.endswith(".png") else "JPEG"
        img = half(120, 180, name.endswith("page1.png"), fmt)
        z.writestr(zipfile.ZipInfo(name), img, compress_type=zipfile.ZIP_STORED if i % 2 == 0 else zipfile.ZIP_DEFLATED)
    z.writestr("ComicInfo.xml", "<?xml version=\"1.0\"?><ComicInfo><Title>Test Comic</Title><Series>Series X</Series>"
               "<Number>3</Number><Writer>Comic Writer</Writer><LanguageISO>lt</LanguageISO></ComicInfo>")
